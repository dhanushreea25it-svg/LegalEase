from io import BytesIO
import html
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


BRAND_NAME = "LegalEase"


def sanitize_text(text: str) -> str:
    if text is None:
        return ""

    text = str(text)

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u00a0": " ",
        "\u2022": "-",
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text


def is_heading(line: str) -> bool:
    line = line.strip()

    if not line:
        return False

    if re.match(r"^\d+[\.\)]\s+", line):
        return True

    if line.isupper() and len(line) < 100:
        return True

    return False


def format_docx(
    content: str,
    document_type: str = "Legal Document",
) -> bytes:

    document = Document()

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    brand = document.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER

    brand_run = brand.add_run(BRAND_NAME)
    brand_run.bold = True
    brand_run.font.name = "Times New Roman"
    brand_run.font.size = Pt(16)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title.add_run(
        sanitize_text(document_type).upper()
    )
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(14)

    document.add_paragraph()

    lines = sanitize_text(content).splitlines()

    for line in lines:

        paragraph = document.add_paragraph()

        if not line.strip():
            continue

        if is_heading(line):

            run = paragraph.add_run(
                line.strip()
            )

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

        else:

            run = paragraph.add_run(
                line.strip()
            )

            run.font.name = "Times New Roman"
            run.font.size = Pt(11)

    footer = section.footer

    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer_paragraph.add_run(
        "LegalEase - AI-generated drafting aid | "
        "Legal review recommended"
    )

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(8)

    output = BytesIO()

    document.save(output)

    return output.getvalue()


class LegalPDF(FPDF):

    def __init__(
        self,
        document_type="Legal Document",
    ):
        super().__init__()

        self.document_type = document_type

    def header(self):

        self.set_font(
            "Helvetica",
            "B",
            15,
        )

        self.cell(
            0,
            10,
            BRAND_NAME,
            align="C",
        )

        self.ln(5)

        self.set_font(
            "Helvetica",
            "B",
            11,
        )

        self.cell(
            0,
            8,
            sanitize_text(
                self.document_type
            ),
            align="C",
        )

        self.ln(8)

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "",
            8,
        )

        self.cell(
            0,
            10,
            "LegalEase - AI-generated drafting aid | "
            "Legal review recommended",
            align="C",
        )


def format_pdf(
    content: str,
    document_type: str = "Legal Document",
) -> bytes:

    pdf = LegalPDF(
        document_type
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "",
        11,
    )

    content = sanitize_text(
        content
    )

    # Remove problematic characters
    content = "".join(
        char
        for char in content
        if char == "\n"
        or char == "\t"
        or char.isprintable()
    )

    max_chars = 85

    for original_line in content.splitlines():

        line = original_line.strip()

        if not line:
            pdf.ln(4)
            continue

        words = line.split()

        current_line = ""

        for word in words:

            test_line = (
                word
                if not current_line
                else current_line + " " + word
            )

            if len(test_line) <= max_chars:

                current_line = test_line

            else:

                if current_line:

                    pdf.set_x(
                        pdf.l_margin
                    )

                    if is_heading(
                        current_line
                    ):

                        pdf.set_font(
                            "Helvetica",
                            "B",
                            11,
                        )

                    else:

                        pdf.set_font(
                            "Helvetica",
                            "",
                            11,
                        )

                    pdf.cell(
                        0,
                        6,
                        current_line,
                        new_x="LMARGIN",
                        new_y="NEXT",
                    )

                current_line = word

        if current_line:

            pdf.set_x(
                pdf.l_margin
            )

            if is_heading(
                current_line
            ):

                pdf.set_font(
                    "Helvetica",
                    "B",
                    11,
                )

            else:

                pdf.set_font(
                    "Helvetica",
                    "",
                    11,
                )

            pdf.cell(
                0,
                6,
                current_line,
                new_x="LMARGIN",
                new_y="NEXT",
            )

        pdf.ln(2)

    return bytes(
        pdf.output()
    )


def format_txt(
    content: str,
) -> bytes:

    return sanitize_text(
        content
    ).encode("utf-8")


def format_html_preview(
    content: str,
) -> str:

    safe_content = html.escape(
        sanitize_text(content)
    )

    safe_content = safe_content.replace(
        "\n",
        "<br>",
    )

    return f"""
    <div class="document-preview">

        <div class="document-brand">
            {BRAND_NAME}
        </div>

        <div class="document-content">
            {safe_content}
        </div>

    </div>
    """